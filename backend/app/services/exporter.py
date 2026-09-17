import io
from datetime import datetime
from typing import List
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from app.models.meeting import Meeting
from app.models.mom import MinutesOfMeeting
from app.models.action_item import ActionItem

def generate_mom_pdf(meeting: Meeting, mom: MinutesOfMeeting, action_items: List[ActionItem]) -> io.BytesIO:
    """
    Generates a professionally styled institutional PDF for Minutes of Meeting.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#4B5563'),
        spaceAfter=15
    )
    
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#1E3A8A'),
        spaceBefore=12,
        spaceAfter=6
    )
    
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#1F2937'),
        spaceAfter=6
    )

    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#1F2937'),
        leftIndent=15,
        spaceAfter=4
    )

    story = []

    # Header / Branding
    story.append(Paragraph("CONVERSEIQ ACADEMIC INTELLIGENCE", ParagraphStyle('SubHeader', fontName='Helvetica-Bold', fontSize=9, textColor=colors.HexColor('#3B82F6'))))
    story.append(Paragraph(f"Minutes of Meeting: {meeting.title}", title_style))
    
    meta_text = f"<b>Date:</b> {meeting.date.strftime('%B %d, %Y %I:%M %p')} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Duration:</b> {meeting.duration_minutes} mins &nbsp;&nbsp;|&nbsp;&nbsp; <b>Status:</b> {meeting.status}"
    story.append(Paragraph(meta_text, subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#E5E7EB'), spaceAfter=15))

    # Participants
    if meeting.participants:
        story.append(Paragraph("Meeting Attendees", h2_style))
        attendees_text = ", ".join([p.name + (f" ({p.speaker_label})" if p.speaker_label else "") for p in meeting.participants])
        story.append(Paragraph(attendees_text, body_style))
        story.append(Spacer(1, 8))

    # Executive Summary
    story.append(Paragraph("Executive Summary", h2_style))
    story.append(Paragraph(mom.summary if mom else "No summary recorded.", body_style))
    story.append(Spacer(1, 8))

    # Key Decisions
    if mom and mom.decisions:
        story.append(Paragraph("Key Decisions Agreed", h2_style))
        for idx, decision in enumerate(mom.decisions, 1):
            story.append(Paragraph(f"<b>{idx}.</b> {decision}", bullet_style))
        story.append(Spacer(1, 8))

    # Topics Discussed
    if mom and mom.topics_discussed:
        story.append(Paragraph("Discussion Topics", h2_style))
        for topic in mom.topics_discussed:
            if isinstance(topic, dict):
                t_name = topic.get("topic", "Topic")
                story.append(Paragraph(f"<b>• {t_name}</b>", body_style))
                for pt in topic.get("points", []):
                    story.append(Paragraph(f"- {pt}", bullet_style))
            elif isinstance(topic, str):
                story.append(Paragraph(f"• {topic}", bullet_style))
        story.append(Spacer(1, 8))

    # Action Items Table
    if action_items:
        story.append(Paragraph("Action Items & Accountability", h2_style))
        
        table_data = [["Task", "Owner", "Priority", "Status", "Due Date"]]
        for item in action_items:
            due_str = item.due_date.strftime('%b %d, %Y') if item.due_date else "TBD"
            table_data.append([
                Paragraph(item.task, body_style),
                item.owner_name,
                item.priority,
                item.status,
                due_str
            ])
            
        t = Table(table_data, colWidths=[200, 100, 65, 75, 90])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('TOPPADDING', (0, 0), (-1, 0), 6),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#F9FAFB'), colors.white])
        ]))
        story.append(t)

    # Finalization footer
    story.append(Spacer(1, 20))
    final_note = f"<i>Document Generated by ConverseIQ Production System on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}. Status: {'Finalized' if (mom and mom.is_finalized) else 'Draft'}.</i>"
    story.append(Paragraph(final_note, ParagraphStyle('DocFooter', fontName='Helvetica-Oblique', fontSize=8, textColor=colors.HexColor('#9CA3AF'))))

    doc.build(story)
    buffer.seek(0)
    return buffer

def generate_mom_docx(meeting: Meeting, mom: MinutesOfMeeting, action_items: List[ActionItem]) -> io.BytesIO:
    """
    Generates an editable Word (.docx) document for Minutes of Meeting.
    """
    doc = Document()
    
    # Title
    title = doc.add_heading(f"Minutes of Meeting: {meeting.title}", level=1)
    title.paragraph_format.space_after = Pt(4)
    
    # Metadata
    meta = doc.add_paragraph()
    meta.add_run(f"Date: {meeting.date.strftime('%B %d, %Y %I:%M %p')}  |  Duration: {meeting.duration_minutes} mins  |  Status: {meeting.status}").italic = True
    meta.paragraph_format.space_after = Pt(14)
    
    # Attendees
    if meeting.participants:
        doc.add_heading("Meeting Attendees", level=2)
        attendees_text = ", ".join([p.name + (f" ({p.speaker_label})" if p.speaker_label else "") for p in meeting.participants])
        doc.add_paragraph(attendees_text)
        
    # Executive Summary
    doc.add_heading("Executive Summary", level=2)
    doc.add_paragraph(mom.summary if mom else "No summary recorded.")
    
    # Decisions
    if mom and mom.decisions:
        doc.add_heading("Key Decisions Agreed", level=2)
        for d in mom.decisions:
            doc.add_paragraph(d, style='List Bullet')
            
    # Topics Discussed
    if mom and mom.topics_discussed:
        doc.add_heading("Discussion Topics", level=2)
        for topic in mom.topics_discussed:
            if isinstance(topic, dict):
                p = doc.add_paragraph()
                p.add_run(topic.get("topic", "Topic")).bold = True
                for pt in topic.get("points", []):
                    doc.add_paragraph(pt, style='List Bullet 2')
            elif isinstance(topic, str):
                doc.add_paragraph(topic, style='List Bullet')
                
    # Action Items Table
    if action_items:
        doc.add_heading("Action Items & Accountability", level=2)
        table = doc.add_table(rows=1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Task'
        hdr_cells[1].text = 'Owner'
        hdr_cells[2].text = 'Priority'
        hdr_cells[3].text = 'Status'
        hdr_cells[4].text = 'Due Date'
        
        for item in action_items:
            row_cells = table.add_row().cells
            row_cells[0].text = item.task
            row_cells[1].text = item.owner_name
            row_cells[2].text = item.priority
            row_cells[3].text = item.status
            row_cells[4].text = item.due_date.strftime('%b %d, %Y') if item.due_date else "TBD"
            
    # Save to buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
